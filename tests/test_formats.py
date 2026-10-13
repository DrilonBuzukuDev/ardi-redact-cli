from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from ard_redact_cli.detection import Detector
from ard_redact_cli.formats import prepare
from ard_redact_cli.processing import process_file


def test_plain_text_preserves_bom_crlf_and_original(tmp_path: Path, detector: Detector) -> None:
  path = tmp_path / "notes.txt"
  original = b"\xef\xbb\xbfJohn Smith\r\nlives in Vienna.\r\nJohn Smith"
  path.write_bytes(original)
  result = process_file(path, detector)
  assert path.read_bytes() == original
  assert result.output.read_bytes() == (
    b"\xef\xbb\xbf[PERSON_1]\r\nlives in [LOCATION_1].\r\n[PERSON_1]"
  )


def test_markdown_spanning_formatting(tmp_path: Path, detector: Detector) -> None:
  path = tmp_path / "notes.md"
  path.write_text(
    "# **John** Smith\n\n[John Smith](https://example.org)\n"
    "[contact](mailto:john_smith@example.com)\n"
    "```text\nJohn Smith\n```\n`john@example.com`\n",
    encoding="utf-8",
  )
  result = process_file(path, detector)
  text = result.output.read_text()
  assert "# **[PERSON_1]**" in text
  assert "[[PERSON_1]](https://example.org)" in text
  assert "[contact](mailto:[EMAIL_1])" in text
  assert "```text\n[PERSON_1]\n```" in text
  assert "`[EMAIL_2]`" in text
  assert "John" not in text and "john" not in text


@pytest.mark.parametrize(
  "suffix,timing,header",
  [
    (".srt", "00:00:01,000 --> 00:00:05,000", ""),
    (".vtt", "00:01.000 --> 00:05.000 align:start", "WEBVTT\r\n\r\n"),
  ],
)
def test_subtitle_multiline_structure(
  tmp_path: Path, detector: Detector, suffix: str, timing: str, header: str
) -> None:
  path = tmp_path / f"interview{suffix}"
  source = f"{header}cue-1\r\n{timing}\r\n<i>John\r\nSmith</i> in Vienna.\r\n\r\n"
  path.write_bytes(source.encode())
  result = process_file(path, detector)
  assert (
    result.output.read_bytes()
    == (f"{header}cue-1\r\n{timing}\r\n<i>[PERSON_1]\r\n</i> in [LOCATION_1].\r\n\r\n").encode()
  )


def test_vtt_metadata_and_cue_id_are_preserved(tmp_path: Path, detector: Detector) -> None:
  path = tmp_path / "notes.vtt"
  path.write_text("WEBVTT\n\nNOTE Alice\nprivate\n\nAlice\n00:01.000 --> 00:03.000\nAlice\n")
  result = process_file(path, detector)
  assert "NOTE Alice\nprivate" in result.output.read_text()
  assert "\nAlice\n00:01.000" in result.output.read_text()
  assert result.counts == {"PERSON": 1}
  assert result.warnings


@pytest.mark.parametrize("suffix,source", [(".srt", "not a cue"), (".vtt", "not WEBVTT")])
def test_malformed_subtitles_fail(tmp_path: Path, suffix: str, source: str) -> None:
  path = tmp_path / f"bad{suffix}"
  path.write_text(source)
  with pytest.raises(ValueError):
    prepare(path)


def test_docx_runs_tables_headers_footers(tmp_path: Path, detector: Detector) -> None:
  path = tmp_path / "letter.docx"
  document = Document()
  paragraph = document.add_paragraph()
  paragraph.add_run("John ").bold = True
  paragraph.add_run("Smith").italic = True
  paragraph.add_run(" lives in Vienna.")
  table = document.add_table(rows=1, cols=1)
  table.cell(0, 0).text = "John Smith"
  nested = table.cell(0, 0).add_table(rows=1, cols=1)
  nested.cell(0, 0).text = "john@example.com"
  document.sections[0].header.paragraphs[0].text = "John Smith"
  document.sections[0].footer.paragraphs[0].text = "Vienna"
  document.save(str(path))
  original = path.read_bytes()
  result = process_file(path, detector)
  updated = Document(str(result.output))
  assert updated.paragraphs[0].text == "[PERSON_1] lives in [LOCATION_1]."
  assert updated.paragraphs[0].runs[0].bold is True
  assert updated.paragraphs[0].runs[1].italic is True
  assert updated.tables[0].cell(0, 0).text == "[PERSON_1]\n"
  assert updated.tables[0].cell(0, 0).tables[0].cell(0, 0).text == "[EMAIL_1]"
  assert updated.sections[0].header.paragraphs[0].text == "[PERSON_1]"
  assert updated.sections[0].footer.paragraphs[0].text == "[LOCATION_1]"
  assert path.read_bytes() == original
  assert result.counts == {"PERSON": 3, "LOCATION": 2, "EMAIL": 1}
  with ZipFile(path) as before, ZipFile(result.output) as after:
    assert before.namelist() == after.namelist()
    for name in before.namelist():
      if name not in {"word/document.xml", "word/header1.xml", "word/footer1.xml"}:
        assert before.read(name) == after.read(name)


def test_docx_breaks_preserved(tmp_path: Path, detector: Detector) -> None:
  path = tmp_path / "break.docx"
  document = Document()
  run = document.add_paragraph().add_run("John")
  run.add_break()
  run.add_text("Smith")
  document.save(str(path))
  result = process_file(path, detector)
  assert Document(str(result.output)).paragraphs[0].text == "[PERSON_1]\n"


@pytest.mark.parametrize("encoding,bom", [("utf-16-le", b"\xff\xfe"), ("utf-16-be", b"\xfe\xff")])
def test_utf16(tmp_path: Path, detector: Detector, encoding: str, bom: bytes) -> None:
  path = tmp_path / "utf16.txt"
  path.write_bytes(bom + "Alice\r\n".encode(encoding))
  result = process_file(path, detector)
  assert result.output.read_bytes() == bom + "[PERSON_1]\r\n".encode(encoding)


def test_docx_hyperlink_text_and_unsupported_content(tmp_path: Path, detector: Detector) -> None:
  path = tmp_path / "hyperlink.docx"
  document = Document()
  document.add_paragraph()
  document.save(str(path))
  # Fixture XML covers structures with no public python-docx creation API.
  fixture = (
    b'<w:p><w:hyperlink><w:r><w:t xml:space="preserve">John </w:t></w:r>'
    b"<w:r><w:t>Smith</w:t></w:r></w:hyperlink>"
    b"<w:r><w:pict><w:txbxContent><w:p><w:r><w:t>Alice</w:t></w:r></w:p>"
    b"</w:txbxContent></w:pict></w:r></w:p>"
  )
  with ZipFile(BytesIO(path.read_bytes())) as original, ZipFile(path, "w") as modified:
    for info in original.infolist():
      content = original.read(info)
      if info.filename == "word/document.xml":
        assert b"<w:p/>" in content
        content = content.replace(b"<w:p/>", fixture, 1)
      modified.writestr(info, content)
  result = process_file(path, detector)
  assert result.counts == {"PERSON": 1}
  assert "DOCX text boxes are not scanned." in result.warnings
  with ZipFile(result.output) as archive:
    xml = archive.read("word/document.xml")
    assert b"[PERSON_1]" in xml and b"Alice" in xml
    assert b"John" not in xml and b"Smith" not in xml
