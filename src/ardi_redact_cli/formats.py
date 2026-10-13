"""Format adapters retain structural bytes or XML and expose only editable text."""

import re
from collections.abc import Iterator
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import BinaryIO, Protocol
from zipfile import ZipFile, ZipInfo

from docx import Document
from lxml import etree
from lxml.etree import _Element as XmlElement  # pyright: ignore[reportPrivateUsage]

from ardi_redact_cli.detection import PLACEHOLDER_PATTERN
from ardi_redact_cli.redaction import TextUnit, apply_edits

SUPPORTED = frozenset({".txt", ".md", ".srt", ".vtt", ".docx"})
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


class PreparedFile(Protocol):
  units: list[TextUnit]
  warnings: list[str]

  def write(self, edits: list[dict[int, str]], stream: BinaryIO) -> None: ...


def _protect_matches(source: str, patterns: list[str]) -> set[int]:
  protected: set[int] = set()
  for pattern in patterns:
    for match in re.finditer(pattern, source, re.MULTILINE):
      protected.update(range(match.start(), match.end()))
  return protected


def markdown_unit(source: str) -> TextUnit:
  # Hide syntax from NER, but keep the content of code blocks, inline code,
  # link labels, and link destinations eligible for detection.
  protected = _protect_matches(
    source,
    [
      r"^[ \t]{0,3}(?:`{3,}|~{3,})[^\r\n]*",  # fence + language
      r"^[ \t]{0,3}(?:#{1,6}\s+|(?:>\s*)+|[-+*]\s+|\d+[.)]\s+)",
      r"^[ \t]*(?:[-*_][ \t]*){3,}$",  # thematic rules
      r"!?\[|\]\(|\](?=\[)|\[|\]|[()]|`+|\*+|(?<!\w)_+|_+(?!\w)|~~|(?<=\n)\|",
      r"\\(?=[!\"#$%&'()*+,\-./:;<=>?@\[\]\\^_`{|}~])",
    ],
  )
  # Retain existing placeholders as tokens so the detector can protect them.
  boundaries = _protect_matches(source, [r"\]\(|\](?=\[)|(?<=\()mailto:"])
  for match in PLACEHOLDER_PATTERN.finditer(source):
    protected.difference_update(range(match.start(), match.end()))
  return TextUnit.project(source, protected, boundaries)


@dataclass
class TextFile:
  chunks: list[str]
  units: list[TextUnit]
  editable: list[int]
  encoding: str
  warnings: list[str]
  bom: bytes = b""

  def write(self, edits: list[dict[int, str]], stream: BinaryIO) -> None:
    chunks = self.chunks.copy()
    for index, changes in zip(self.editable, edits, strict=True):
      chunks[index] = apply_edits(chunks[index], changes)
    stream.write(self.bom + "".join(chunks).encode(self.encoding))


TIMING = re.compile(
  r"^(?:\d{2,}:)?\d{2}:\d{2}[.,]\d{3}\s+-->\s+"
  r"(?:\d{2,}:)?\d{2}:\d{2}[.,]\d{3}(?:[ \t]+[^\r\n]*)?[\r\n]*$"
)


def _subtitle(source: str, suffix: str, encoding: str, bom: bytes = b"") -> TextFile:
  blocks = re.split(r"((?:\r?\n){2,})", source)
  units: list[TextUnit] = []
  chunks: list[str] = []
  editable: list[int] = []
  warnings: list[str] = []
  if suffix == ".vtt" and not source.startswith("WEBVTT"):
    raise ValueError("VTT input must start with WEBVTT.")
  for block in blocks:
    lines = block.splitlines(keepends=True)
    if not lines or not block.strip():
      chunks.append(block)
      continue
    first = lines[0].strip()
    if suffix == ".vtt" and (
      first.startswith("WEBVTT")
      or first in {"STYLE", "REGION"}
      or first == "NOTE"
      or first.startswith("NOTE ")
      or first.startswith("NOTE\t")
    ):
      chunks.append(block)
      warnings.append("VTT metadata, notes, and styles are preserved and not scanned.")
      continue
    timing = next((index for index, line in enumerate(lines[:2]) if TIMING.match(line)), None)
    if timing is None:
      raise ValueError("Invalid subtitle cue: expected an identifier and timestamp line.")
    chunks.append("".join(lines[: timing + 1]))
    content = "".join(lines[timing + 1 :])
    editable.append(len(chunks))
    chunks.append(content)
    protected = _protect_matches(content, [r"<[^>]*>", r"\{\\[^}]*\}"])
    units.append(TextUnit.project(content, protected))
  return TextFile(chunks, units, editable, encoding, sorted(set(warnings)), bom)


@dataclass
class XmlParagraph:
  nodes: list[tuple[XmlElement, int, str]]


class DocxFile:
  def __init__(self, data: bytes) -> None:
    # python-docx checks the Word package. Modify text nodes rather than
    # assigning Run.text, which would discard drawings and other run content.
    Document(BytesIO(data))
    self.units: list[TextUnit] = []
    self.warnings: list[str] = []
    self.paragraphs: list[XmlParagraph] = []
    self.entries: list[tuple[ZipInfo, bytes]] = []
    self.roots: dict[str, XmlElement] = {}
    with ZipFile(BytesIO(data)) as archive:
      self.comment = archive.comment
      for info in archive.infolist():
        content = archive.read(info)
        self.entries.append((info, content))
        name = info.filename
        if name == "word/document.xml" or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name):
          root = etree.fromstring(content, etree.XMLParser(resolve_entities=False))
          self.roots[name] = root
          self._paragraphs(root)
        elif name in {"word/comments.xml", "word/footnotes.xml", "word/endnotes.xml"}:
          self.warnings.append("DOCX comments, footnotes, and endnotes are not scanned.")
        elif name.startswith(("word/media/", "word/embeddings/")):
          self.warnings.append("DOCX images and embedded objects are not scanned.")
        elif name.endswith(".rels") and b'TargetMode="External"' in content:
          self.warnings.append("DOCX external hyperlink targets are not scanned.")
    self.warnings.append("DOCX document properties and other package metadata are not scanned.")
    self.warnings = sorted(set(self.warnings))

  def _paragraphs(self, root: XmlElement) -> None:
    if next(root.iter(W + "txbxContent"), None) is not None:
      self.warnings.append("DOCX text boxes are not scanned.")
    if next(root.iter(W + "instrText"), None) is not None:
      self.warnings.append("DOCX field instructions are not scanned; displayed text is scanned.")
    if next(root.iter(W + "del"), None) is not None:
      self.warnings.append("DOCX deleted tracked changes are not scanned.")
    if next(root.iter(W + "altChunk"), None) is not None:
      self.warnings.append("DOCX imported alternative-format content is not scanned.")
    for paragraph in root.iter(W + "p"):
      if any(parent.tag == W + "txbxContent" for parent in paragraph.iterancestors()):
        continue
      source = ""
      nodes: list[tuple[XmlElement, int, str]] = []
      protected: set[int] = set()
      for node in self._text_nodes(paragraph):
        if node.tag == W + "t":
          value = node.text or ""
          nodes.append((node, len(source), value))
          source += value
        else:
          protected.add(len(source))
          source += "\n" if node.tag in {W + "br", W + "cr"} else "\t"
      if nodes:
        # Tabs and line breaks are represented by whitespace for NER.
        unit = TextUnit.project(source)
        unit.offsets = [None if index in protected else index for index in unit.offsets]
        self.units.append(unit)
        self.paragraphs.append(XmlParagraph(nodes))

  def _text_nodes(self, paragraph: XmlElement) -> Iterator[XmlElement]:
    for node in paragraph.iter():
      if node.tag not in {W + "t", W + "tab", W + "br", W + "cr"}:
        continue
      # Nested text boxes have their own paragraphs and are unsupported.
      if any(parent.tag in {W + "txbxContent", W + "del"} for parent in node.iterancestors()):
        continue
      yield node

  def write(self, edits: list[dict[int, str]], stream: BinaryIO) -> None:
    for paragraph, changes in zip(self.paragraphs, edits, strict=True):
      for node, start, original in paragraph.nodes:
        updated = apply_edits(original, changes, start)
        node.text = updated
        if updated.startswith(" ") or updated.endswith(" "):
          node.set(XML_SPACE, "preserve")
    with ZipFile(stream, "w") as archive:
      archive.comment = self.comment
      for info, content in self.entries:
        if info.filename in self.roots:
          content = etree.tostring(
            self.roots[info.filename],
            encoding="UTF-8",
            xml_declaration=True,
            standalone=True,
          )
        archive.writestr(info, content)


def prepare(path: Path) -> PreparedFile:
  suffix = path.suffix.lower()
  if suffix not in SUPPORTED:
    raise ValueError("Unsupported file format.")
  data = path.read_bytes()
  if suffix == ".docx":
    return DocxFile(data)
  bom = b""
  encoding = "utf-8"
  if data.startswith((b"\xff\xfe", b"\xfe\xff")):
    bom, data = data[:2], data[2:]
    encoding = "utf-16-le" if bom == b"\xff\xfe" else "utf-16-be"
  if data.startswith(b"\xef\xbb\xbf"):
    encoding = "utf-8-sig"
  source = data.decode(encoding)
  if suffix in {".vtt", ".srt"}:
    return _subtitle(source, suffix, encoding, bom)
  unit = markdown_unit(source) if suffix == ".md" else TextUnit.project(source)
  return TextFile([source], [unit], [0], encoding, [], bom)
